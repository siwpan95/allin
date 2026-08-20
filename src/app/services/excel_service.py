# -*- coding: utf-8 -*-
"""
Excel 配置导入服务
支持从 .xlsx 文件读取游戏配置，方便策划通过 Excel 更新游戏内容。
Excel 文件优先于 JSON 文件加载。
"""
import os
import sys
from typing import Dict, List, Optional

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ExcelService:
    """Excel 配置服务 — 读取 .xlsx 配置文件，返回与 JSON 一致的数据结构"""

    # 各 Sheet 的列映射
    ENEMY_COLS = ["名称", "等级", "HP", "伤害", "防御", "经验", "金币", "掉落表", "种族", "描述"]
    QUEST_COLS = ["任务ID", "任务名称", "类型", "描述", "目标", "目标数量",
                  "等级需求", "前置任务", "奖励经验", "奖励金币", "是否史诗", "故事文本", "完成文本"]
    WEAPON_TYPE_NAMES = ["SWORD", "AXE", "MACE", "DAGGER", "STAFF", "BOW", "CROSSBOW", "POLEARM"]
    ARMOR_TYPE_NAMES = ["LIGHT", "MEDIUM", "HEAVY", "ROBE"]
    QUALITY_NAMES = ["CRUDE", "NORMAL", "FINE", "SUPERIOR", "EXCELLENT", "MASTER", "LEGENDARY", "ARTIFACT"]
    POTION_COLS = ["药水名称", "描述", "持续时间(秒)", "掉落权重"]
    ARTIFACT_COLS = ["神器ID", "名称", "描述", "效果类型", "效果值", "来源"]
    ACHIEVEMENT_COLS = ["成就ID", "名称", "描述", "目标值", "奖励类型", "奖励值", "隐藏"]

    def __init__(self, excel_path: str = None):
        """初始化 Excel 服务

        Args:
            excel_path: Excel 文件路径，默认与 JSON 同目录
        """
        if not HAS_OPENPYXL:
            raise ImportError(
                "openpyxl 未安装。请运行: pip install openpyxl\n"
                "或添加到 requirements.txt: openpyxl>=3.1.0"
            )

        if excel_path is None:
            if getattr(sys, 'frozen', False):
                # 打包为 exe 后，从 exe 同目录下的 resources/config/ 读取外部 Excel 配置
                base_dir = os.path.dirname(sys.executable)
                excel_path = os.path.join(base_dir, "resources", "config", "game_config.xlsx")
            else:
                # 开发环境：从源码目录读取
                current_dir = os.path.dirname(os.path.abspath(__file__))
                src_dir = os.path.dirname(os.path.dirname(current_dir))
                config_dir = os.path.join(src_dir, "resources", "config")
                excel_path = os.path.join(config_dir, "game_config.xlsx")
        self.excel_path = excel_path
        self._cache = {}

    # ======================== 核心读取方法 ========================

    def file_exists(self) -> bool:
        """检查 Excel 配置文件是否存在"""
        return os.path.isfile(self.excel_path)

    def _open_workbook(self) -> Workbook:
        """打开工作簿（带缓存）"""
        if "_wb" in self._cache:
            return self._cache["_wb"]
        if not self.file_exists():
            raise FileNotFoundError(f"Excel 配置文件不存在: {self.excel_path}")
        wb = load_workbook(self.excel_path, read_only=True, data_only=True)
        self._cache["_wb"] = wb
        return wb

    def clear_cache(self):
        """清除缓存，强制重新读取 Excel"""
        if "_wb" in self._cache:
            self._cache["_wb"].close()
        self._cache.clear()

    @staticmethod
    def _parse_bool(value) -> bool:
        """将 Excel 中的值解析为布尔类型"""
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            return value.strip().lower() in ("true", "是", "yes", "1")
        if isinstance(value, (int, float)):
            return bool(value)
        return False

    @staticmethod
    def _parse_float(value) -> float:
        """安全解析为浮点数"""
        try:
            return float(value) if value is not None else 0.0
        except (ValueError, TypeError):
            return 0.0

    @staticmethod
    def _parse_int(value) -> int:
        """安全解析为整数"""
        try:
            return int(value) if value is not None else 0
        except (ValueError, TypeError):
            return 0

    def _read_sheet_to_dicts(self, sheet_name: str, columns: List[str],
                             start_row: int = 2) -> List[dict]:
        """通用方法：将 Sheet 数据读取为字典列表

        Args:
            sheet_name: Sheet 名称
            columns:  列名列表
            start_row: 数据起始行（默认第2行，第1行为表头）

        Returns:
            字典列表
        """
        wb = self._open_workbook()
        if sheet_name not in wb.sheetnames:
            print(f"警告: Sheet '{sheet_name}' 不存在")
            return []

        ws = wb[sheet_name]
        result = []
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            if row is None or all(c is None for c in row):
                continue  # 跳过空行
            row_dict = {}
            for i, col_name in enumerate(columns):
                if i < len(row):
                    row_dict[col_name] = row[i]
                else:
                    row_dict[col_name] = None
            result.append(row_dict)
        return result

    # ======================== 敌人模板 ========================

    def get_enemy_templates(self) -> List[dict]:
        """读取敌人模板 → 与 JSON enemy_templates 一致的格式"""
        raw = self._read_sheet_to_dicts("敌人模板", self.ENEMY_COLS)
        result = []
        for r in raw:
            result.append({
                "name": str(r.get("名称", "")).strip(),
                "level": self._parse_int(r.get("等级")),
                "hp": self._parse_int(r.get("HP")),
                "damage": self._parse_int(r.get("伤害")),
                "defense": self._parse_int(r.get("防御")),
                "exp": self._parse_int(r.get("经验")),
                "gold": self._parse_int(r.get("金币")),
                "loot_table": str(r.get("掉落表", "normal")).strip().lower(),
                "race": str(r.get("种族", "")).strip(),
                "description": str(r.get("描述", "")).strip(),
            })
        return result

    # ======================== 任务模板 ========================

    def get_quest_templates(self) -> List[dict]:
        """读取任务模板 → 与 JSON quest_templates 一致的格式"""
        raw = self._read_sheet_to_dicts("任务模板", self.QUEST_COLS)
        result = []
        for r in raw:
            quest = {
                "quest_id": str(r.get("任务ID", "")).strip(),
                "name": str(r.get("任务名称", "")).strip(),
                "type": str(r.get("类型", "")).strip().upper(),
                "description": str(r.get("描述", "")).strip(),
                "target": str(r.get("目标", "")).strip(),
                "target_count": self._parse_int(r.get("目标数量")),
                "level_req": self._parse_int(r.get("等级需求")),
                "reward_exp": self._parse_int(r.get("奖励经验")),
                "reward_gold": self._parse_int(r.get("奖励金币")),
                "story_text": str(r.get("故事文本", "")).strip(),
                "completion_text": str(r.get("完成文本", "")).strip(),
            }
            prev = str(r.get("前置任务", "")).strip()
            if prev:
                quest["previous_quest"] = prev
            is_epic = self._parse_bool(r.get("是否史诗"))
            if is_epic:
                quest["is_epic"] = True
            result.append(quest)
        return result

    # ======================== 物品模板 ========================

    def get_item_templates(self) -> dict:
        """读取物品模板 → 与 JSON item_templates 一致的格式"""
        return {
            "weapon_templates": self._read_weapon_templates(),
            "armor_templates": self._read_armor_templates(),
            "jewelry_names": self._read_jewelry_names(),
            "quality_settings": self._read_quality_settings(),
            "prefixes": self._read_prefixes(),
            "effect_pool": self._read_effect_pool(),
        }

    def _read_weapon_templates(self) -> dict:
        """读取武器模板 Sheet"""
        wb = self._open_workbook()
        if "武器模板" not in wb.sheetnames:
            return {}
        ws = wb["武器模板"]
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            weapon_type = str(row[0]).strip() if row[0] else ""
            if not weapon_type:
                continue
            names = [str(row[i]).strip() for i in range(1, 6) if i < len(row) and row[i]]
            result[weapon_type] = names
        return result

    def _read_armor_templates(self) -> dict:
        """读取护甲模板 Sheet"""
        wb = self._open_workbook()
        if "护甲模板" not in wb.sheetnames:
            return {}
        ws = wb["护甲模板"]
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            armor_type = str(row[0]).strip() if row[0] else ""
            if not armor_type:
                continue
            names = [str(row[i]).strip() for i in range(1, 6) if i < len(row) and row[i]]
            result[armor_type] = names
        return result

    def _read_jewelry_names(self) -> dict:
        """读取饰品名称 Sheet"""
        wb = self._open_workbook()
        if "饰品名称" not in wb.sheetnames:
            return {"NECK": [], "RING": []}
        ws = wb["饰品名称"]
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            jewelry_type = str(row[0]).strip() if row[0] else ""
            if not jewelry_type:
                continue
            names = [str(row[i]).strip() for i in range(1, 6) if i < len(row) and row[i]]
            result[jewelry_type] = names
        return result

    def _read_quality_settings(self) -> dict:
        """读取品质设置 Sheet"""
        wb = self._open_workbook()
        if "品质设置" not in wb.sheetnames:
            return {"QUALITY_DAMAGE_MULT": {}, "QUALITY_DEFENSE_MULT": {}}
        ws = wb["品质设置"]
        damage_mult = {}
        defense_mult = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            quality = str(row[0]).strip().upper() if row[0] else ""
            if not quality:
                continue
            if len(row) >= 2:
                damage_mult[quality] = self._parse_float(row[1])
            if len(row) >= 3:
                defense_mult[quality] = self._parse_float(row[2])
        return {
            "QUALITY_DAMAGE_MULT": damage_mult,
            "QUALITY_DEFENSE_MULT": defense_mult,
        }

    def _read_prefixes(self) -> dict:
        """读取前缀词缀 Sheet"""
        wb = self._open_workbook()
        if "前缀词缀" not in wb.sheetnames:
            return {"weapon": [], "armor": [], "accessory": []}
        ws = wb["前缀词缀"]
        result = {"weapon": [], "armor": [], "accessory": []}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            prefix_type = str(row[0]).strip() if row[0] else ""
            prefix_text = str(row[1]).strip() if len(row) >= 2 and row[1] else ""
            if prefix_type and prefix_text:
                if prefix_type in result:
                    result[prefix_type].append(prefix_text)
        return result

    def _read_effect_pool(self) -> list:
        """读取特效池 Sheet"""
        wb = self._open_workbook()
        if "特效池" not in wb.sheetnames:
            return []
        ws = wb["特效池"]
        effects = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or row[0] is None:
                continue
            effects.append(str(row[0]).strip())
        return effects

    # ======================== 药水配置 ========================

    def get_potion_config(self) -> List[dict]:
        """读取药水配置"""
        raw = self._read_sheet_to_dicts("药水配置", self.POTION_COLS)
        if not raw:
            return []
        result = []
        for r in raw:
            result.append({
                "name": str(r.get("药水名称", "")).strip(),
                "description": str(r.get("描述", "")).strip(),
                "duration": self._parse_float(r.get("持续时间(秒)")),
                "weight": self._parse_int(r.get("掉落权重")),
            })
        return result

    # ======================== 神器配置 ========================

    def get_artifact_config(self) -> List[dict]:
        """读取神器配置"""
        raw = self._read_sheet_to_dicts("神器配置", self.ARTIFACT_COLS)
        if not raw:
            return []
        result = []
        for r in raw:
            result.append({
                "artifact_id": str(r.get("神器ID", "")).strip(),
                "name": str(r.get("名称", "")).strip(),
                "description": str(r.get("描述", "")).strip(),
                "effect_type": str(r.get("效果类型", "")).strip(),
                "effect_value": self._parse_float(r.get("效果值")),
                "source": str(r.get("来源", "")).strip(),
            })
        return result

    # ======================== 成就配置 ========================

    def get_achievement_config(self) -> List[dict]:
        """读取成就配置"""
        raw = self._read_sheet_to_dicts("成就配置", self.ACHIEVEMENT_COLS)
        if not raw:
            return []
        result = []
        for r in raw:
            result.append({
                "ach_id": str(r.get("成就ID", "")).strip(),
                "name": str(r.get("名称", "")).strip(),
                "description": str(r.get("描述", "")).strip(),
                "target": self._parse_int(r.get("目标值")),
                "reward_type": str(r.get("奖励类型", "")).strip(),
                "reward_value": self._parse_float(r.get("奖励值")),
                "hidden": self._parse_bool(r.get("隐藏")),
            })
        return result

    # ======================== 便捷查询方法 ========================

    def get_enemy_by_name(self, name: str) -> Optional[dict]:
        """根据名称获取敌人模板"""
        for enemy in self.get_enemy_templates():
            if enemy["name"] == name:
                return enemy
        return None

    def get_enemies_by_level(self, level: int, level_range: int = 3) -> List[dict]:
        """根据等级获取合适的敌人"""
        enemies = self.get_enemy_templates()
        suitable = [e for e in enemies if abs(e["level"] - level) <= level_range]
        return suitable if suitable else enemies

    # ======================== 导出 Excel 方法 ========================

    @classmethod
    def export_from_json(cls, json_data_dir: str, output_path: str):
        """从 JSON 文件生成 Excel 模板（策划操作入口）"""
        import json

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # 移除默认 sheet

        # --- Sheet 1: 敌人模板 ---
        ws_enemy = wb.create_sheet("敌人模板")
        for i, col in enumerate(cls.ENEMY_COLS, 1):
            ws_enemy.cell(row=1, column=i, value=col).font = openpyxl.styles.Font(bold=True)

        enemy_file = os.path.join(json_data_dir, "enemy_templates.json")
        if os.path.isfile(enemy_file):
            with open(enemy_file, "r", encoding="utf-8") as f:
                enemy_data = json.load(f)
            enemies = enemy_data.get("enemy_templates", [])
            for r_idx, e in enumerate(enemies, 2):
                ws_enemy.cell(row=r_idx, column=1, value=e.get("name", ""))
                ws_enemy.cell(row=r_idx, column=2, value=e.get("level", 0))
                ws_enemy.cell(row=r_idx, column=3, value=e.get("hp", 0))
                ws_enemy.cell(row=r_idx, column=4, value=e.get("damage", 0))
                ws_enemy.cell(row=r_idx, column=5, value=e.get("defense", 0))
                ws_enemy.cell(row=r_idx, column=6, value=e.get("exp", 0))
                ws_enemy.cell(row=r_idx, column=7, value=e.get("gold", 0))
                ws_enemy.cell(row=r_idx, column=8, value=e.get("loot_table", "normal"))
                ws_enemy.cell(row=r_idx, column=9, value=e.get("race", ""))
                ws_enemy.cell(row=r_idx, column=10, value=e.get("description", ""))

        # --- Sheet 2: 任务模板 ---
        ws_quest = wb.create_sheet("任务模板")
        for i, col in enumerate(cls.QUEST_COLS, 1):
            ws_quest.cell(row=1, column=i, value=col).font = openpyxl.styles.Font(bold=True)

        quest_file = os.path.join(json_data_dir, "quest_templates.json")
        if os.path.isfile(quest_file):
            with open(quest_file, "r", encoding="utf-8") as f:
                quest_data = json.load(f)
            quests = quest_data.get("quest_templates", [])
            for r_idx, q in enumerate(quests, 2):
                ws_quest.cell(row=r_idx, column=1, value=q.get("quest_id", ""))
                ws_quest.cell(row=r_idx, column=2, value=q.get("name", ""))
                ws_quest.cell(row=r_idx, column=3, value=q.get("type", ""))
                ws_quest.cell(row=r_idx, column=4, value=q.get("description", ""))
                ws_quest.cell(row=r_idx, column=5, value=q.get("target", ""))
                ws_quest.cell(row=r_idx, column=6, value=q.get("target_count", 0))
                ws_quest.cell(row=r_idx, column=7, value=q.get("level_req", 1))
                ws_quest.cell(row=r_idx, column=8, value=q.get("previous_quest", ""))
                ws_quest.cell(row=r_idx, column=9, value=q.get("reward_exp", 0))
                ws_quest.cell(row=r_idx, column=10, value=q.get("reward_gold", 0))
                is_epic = q.get("is_epic", False)
                ws_quest.cell(row=r_idx, column=11, value="是" if is_epic else "否")
                ws_quest.cell(row=r_idx, column=12, value=q.get("story_text", ""))
                ws_quest.cell(row=r_idx, column=13, value=q.get("completion_text", ""))

        # --- Sheet 3: 武器模板 ---
        ws_weapon = wb.create_sheet("武器模板")
        ws_weapon.cell(row=1, column=1, value="武器类型").font = openpyxl.styles.Font(bold=True)
        for i, name in enumerate(["粗劣", "普通", "精良", "卓越", "神器"], 2):
            ws_weapon.cell(row=1, column=i, value=name).font = openpyxl.styles.Font(bold=True)

        item_file = os.path.join(json_data_dir, "item_templates.json")
        if os.path.isfile(item_file):
            with open(item_file, "r", encoding="utf-8") as f:
                item_data = json.load(f)
            weapons = item_data.get("weapon_templates", {})
            for r_idx, (wtype, names) in enumerate(weapons.items(), 2):
                ws_weapon.cell(row=r_idx, column=1, value=wtype)
                for i, n in enumerate(names[:5], 2):
                    ws_weapon.cell(row=r_idx, column=i, value=n)

        # --- Sheet 4: 护甲模板 ---
        ws_armor = wb.create_sheet("护甲模板")
        ws_armor.cell(row=1, column=1, value="护甲类型").font = openpyxl.styles.Font(bold=True)
        for i, name in enumerate(["粗劣", "普通", "精良", "卓越", "神器"], 2):
            ws_armor.cell(row=1, column=i, value=name).font = openpyxl.styles.Font(bold=True)

        if os.path.isfile(item_file):
            armors = item_data.get("armor_templates", {})
            for r_idx, (atype, names) in enumerate(armors.items(), 2):
                ws_armor.cell(row=r_idx, column=1, value=atype)
                for i, n in enumerate(names[:5], 2):
                    ws_armor.cell(row=r_idx, column=i, value=n)

        # --- Sheet 5: 饰品名称 ---
        ws_jewelry = wb.create_sheet("饰品名称")
        ws_jewelry.cell(row=1, column=1, value="饰品类型").font = openpyxl.styles.Font(bold=True)
        for i, name in enumerate(["粗劣", "普通", "精良", "卓越", "神器"], 2):
            ws_jewelry.cell(row=1, column=i, value=name).font = openpyxl.styles.Font(bold=True)

        if os.path.isfile(item_file):
            jewelry = item_data.get("jewelry_names", {})
            for r_idx, (jtype, names) in enumerate(jewelry.items(), 2):
                ws_jewelry.cell(row=r_idx, column=1, value=jtype)
                for i, n in enumerate(names[:5], 2):
                    ws_jewelry.cell(row=r_idx, column=i, value=n)

        # --- Sheet 6: 品质设置 ---
        ws_quality = wb.create_sheet("品质设置")
        ws_quality.cell(row=1, column=1, value="品质").font = openpyxl.styles.Font(bold=True)
        ws_quality.cell(row=1, column=2, value="伤害倍率").font = openpyxl.styles.Font(bold=True)
        ws_quality.cell(row=1, column=3, value="防御倍率").font = openpyxl.styles.Font(bold=True)

        if os.path.isfile(item_file):
            quality = item_data.get("quality_settings", {})
            dmg_mult = quality.get("QUALITY_DAMAGE_MULT", {})
            def_mult = quality.get("QUALITY_DEFENSE_MULT", {})
            for r_idx, qname in enumerate(cls.QUALITY_NAMES, 2):
                ws_quality.cell(row=r_idx, column=1, value=qname)
                ws_quality.cell(row=r_idx, column=2, value=dmg_mult.get(qname, 1.0))
                ws_quality.cell(row=r_idx, column=3, value=def_mult.get(qname, 1.0))

        # --- Sheet 7: 前缀词缀 ---
        ws_prefix = wb.create_sheet("前缀词缀")
        ws_prefix.cell(row=1, column=1, value="类型").font = openpyxl.styles.Font(bold=True)
        ws_prefix.cell(row=1, column=2, value="前缀").font = openpyxl.styles.Font(bold=True)

        if os.path.isfile(item_file):
            prefixes = item_data.get("prefixes", {})
            r = 2
            for ptype, pvalues in prefixes.items():
                for pv in pvalues:
                    ws_prefix.cell(row=r, column=1, value=ptype)
                    ws_prefix.cell(row=r, column=2, value=pv)
                    r += 1

        # --- Sheet 8: 特效池 ---
        ws_effect = wb.create_sheet("特效池")
        ws_effect.cell(row=1, column=1, value="特效名称").font = openpyxl.styles.Font(bold=True)

        if os.path.isfile(item_file):
            effects = item_data.get("effect_pool", [])
            for r_idx, ef in enumerate(effects, 2):
                ws_effect.cell(row=r_idx, column=1, value=ef)

        # --- Sheet 9: 药水配置 ---
        ws_potion = wb.create_sheet("药水配置")
        for i, col in enumerate(cls.POTION_COLS, 1):
            ws_potion.cell(row=1, column=i, value=col).font = openpyxl.styles.Font(bold=True)

        # 从 potion.py 硬编码值导出
        try:
            from ..models.potion import POTION_DROP_WEIGHTS
            potion_details = {
                "DOUBLE_GOLD":   ("双倍金币",   "怪物多掉一堆金币", 120.0, 15),
                "DOUBLE_ITEM":   ("双倍物品",   "怪物多掉一件装备", 120.0, 15),
                "FRAIL_MONSTER": ("脆弱怪物",   "怪物属性降低50%",   90.0, 10),
                "DOCILE_MONSTER":("温顺怪物",   "怪物几乎不还手",   90.0,  8),
                "ITEM_VALUE":    ("物品价值提升","装备品质权重上移", 120.0, 12),
                "GUARANTEED_DROP":("必掉装备",  "每只怪物必掉装备",  90.0,  8),
                "MORE_MONSTERS": ("更多怪物",   "每房间怪物+5",    120.0,  5),
                "FREE_CAST":     ("免费施法",   "技能不消耗精神力",  90.0,  5),
                "SPEED_EXPLORE": ("加速探索",   "房间推进间隔减半", 120.0, 12),
                "MORE_EXP":      ("更多经验",   "经验获取+50%",    120.0, 10),
            }
            for r_idx, (key, weight) in enumerate(POTION_DROP_WEIGHTS.items(), 2):
                name, desc, dur, _w = potion_details.get(key.name, (key.name, "", 0.0, 0))
                ws_potion.cell(row=r_idx, column=1, value=name)
                ws_potion.cell(row=r_idx, column=2, value=desc)
                ws_potion.cell(row=r_idx, column=3, value=dur)
                ws_potion.cell(row=r_idx, column=4, value=weight)
        except ImportError:
            pass

        # --- Sheet 10: 神器配置 ---
        ws_art = wb.create_sheet("神器配置")
        for i, col in enumerate(cls.ARTIFACT_COLS, 1):
            ws_art.cell(row=1, column=i, value=col).font = openpyxl.styles.Font(bold=True)

        try:
            from ..services.artifact_service import ARTIFACT_DEFS
            for r_idx, (aid, name, desc, etype, evalue, src) in enumerate(ARTIFACT_DEFS, 2):
                ws_art.cell(row=r_idx, column=1, value=aid)
                ws_art.cell(row=r_idx, column=2, value=name)
                ws_art.cell(row=r_idx, column=3, value=desc)
                ws_art.cell(row=r_idx, column=4, value=etype)
                ws_art.cell(row=r_idx, column=5, value=evalue)
                ws_art.cell(row=r_idx, column=6, value=src)
        except ImportError:
            pass

        # --- Sheet 11: 成就配置 ---
        ws_ach = wb.create_sheet("成就配置")
        for i, col in enumerate(cls.ACHIEVEMENT_COLS, 1):
            ws_ach.cell(row=1, column=i, value=col).font = openpyxl.styles.Font(bold=True)

        try:
            from ..services.achievement_service import ACHIEVEMENT_DEFS
            for r_idx, (aid, name, desc, target, rtype, rval, hidden) in enumerate(ACHIEVEMENT_DEFS, 2):
                ws_ach.cell(row=r_idx, column=1, value=aid)
                ws_ach.cell(row=r_idx, column=2, value=name)
                ws_ach.cell(row=r_idx, column=3, value=desc)
                ws_ach.cell(row=r_idx, column=4, value=target)
                ws_ach.cell(row=r_idx, column=5, value=rtype)
                ws_ach.cell(row=r_idx, column=6, value=rval)
                ws_ach.cell(row=r_idx, column=7, value="是" if hidden else "否")
        except ImportError:
            pass

        # 调整列宽
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)) * 2)
                ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
        print(f"Excel 配置文件已生成: {output_path}")
        print(f"包含 {len(wb.sheetnames)} 个工作表: {', '.join(wb.sheetnames)}")
